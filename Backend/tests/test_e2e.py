"""End-to-end API checks: access control, risk tiers, repeat-payment history, uploads, exports, alerts, training jobs."""

import io
import json
import time
import uuid
from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services import datasets
from app.services.scoring import service

c = TestClient(app)
SAMPLES = Path(__file__).resolve().parents[2] / "samples"
pytestmark = pytest.mark.skipif(not service.ready, reason="model not trained")


def make_user():
    email = f"e2e{uuid.uuid4().hex[:8]}@example.com"
    r = c.post("/api/auth/register", json={"name": "E2E User", "email": email, "password": "supersecret1"})
    assert r.status_code == 200
    return {"Authorization": f"Bearer {r.json()['token']}"}


@pytest.fixture(scope="module")
def h():
    return make_user()


def score(h, **kw):
    body = {"amount_inr": 500, "merchant_category": "merchant", **kw}
    r = c.post("/api/transactions/score", headers=h, json=body)
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------- access control
def test_health_is_public():
    assert c.get("/api/health").json()["model_ready"] is True


@pytest.mark.parametrize("path", ["/api/overview", "/api/transactions", "/api/datasets", "/api/model/report", "/api/federated/status",
                                  "/api/papers", "/api/export/flagged?format=csv", "/api/auth/me"])
def test_protected_endpoints_need_login(path):
    assert c.get(path).status_code == 401


# ---------------------------------------------------------------- risk tiers
@pytest.mark.parametrize("category", ["p2p", "merchant", "investment"])
@pytest.mark.parametrize("amount,decision", [(500, "ALLOW"), (69_999, "ALLOW"), (70_000, "STEP_UP_AUTH"), (85_000, "STEP_UP_AUTH"),
                                             (99_999, "STEP_UP_AUTH"), (100_000, "BLOCK"), (150_000, "BLOCK"), (199_000, "BLOCK")])
def test_amount_tiers(h, category, amount, decision):
    assert score(h, amount_inr=amount, merchant_category=category)["decision"] == decision


def test_risk_percent_matches_tier(h):
    low, mod, high = score(h, amount_inr=500), score(h, amount_inr=80_000), score(h, amount_inr=150_000)
    assert 0 <= low["risk_score"] < 40 and 40 <= mod["risk_score"] < 80 and 80 <= high["risk_score"] <= 99
    assert "high_amount" in high["rules"] and "moderate_amount" in mod["rules"]


def test_model_alone_never_makes_high_below_one_lakh(h):
    scam = score(h, amount_inr=48_000, amount_to_average_ratio=6, is_new_device=1, is_new_beneficiary=1, failed_login_attempts_24h=4, transaction_hour=3)
    assert scam["decision"] == "STEP_UP_AUTH" and scam["model_score"] > 50 and scam["risk_level"] == "MODERATE"


@pytest.mark.parametrize("bad", [{"merchant_category": "food"}, {"merchant_category": "utilities"}, {"amount_inr": -1}, {"amount_inr": 0},
                                 {"transaction_hour": 25}, {"is_new_device": 2}, {"amount_inr": "abc"}])
def test_invalid_payment_rejected(h, bad):
    r = c.post("/api/transactions/score", headers=h, json={"amount_inr": 500, "merchant_category": "merchant", **bad})
    assert r.status_code == 422


# ---------------------------------------------------------------- repeat payments from stored history
def test_repeat_to_same_receiver_is_flagged(h):
    s, r = f"S{uuid.uuid4().hex[:6]}", f"R{uuid.uuid4().hex[:6]}"
    out = [score(h, amount_inr=24_500, amount_to_average_ratio=2, sender_id=s, receiver_id=r, merchant_category="p2p") for _ in range(4)]
    assert out[0]["decision"] == "ALLOW"
    assert out[2]["decision"] == "STEP_UP_AUTH" and "repeat_receiver" in out[2]["rules"]
    assert any("same receiver" in x["reason"] for x in out[3]["reasons"])
    other = score(h, amount_inr=24_500, sender_id=s, receiver_id="R-different", merchant_category="p2p")
    assert "repeat_receiver" not in other["rules"]


def test_burst_of_payments_is_flagged(h):
    s = f"B{uuid.uuid4().hex[:6]}"
    res = [score(h, amount_inr=9_000, sender_id=s, receiver_id=f"R{i}", is_new_beneficiary=1, merchant_category="p2p") for i in range(6)]
    assert "burst" in res[-1]["rules"] and res[-1]["decision"] != "ALLOW"


def test_two_payments_to_a_friend_stay_low(h):
    s, r = f"F{uuid.uuid4().hex[:6]}", f"R{uuid.uuid4().hex[:6]}"
    a = score(h, amount_inr=800, sender_id=s, receiver_id=r, merchant_category="p2p")
    b = score(h, amount_inr=650, sender_id=s, receiver_id=r, merchant_category="p2p")
    assert a["decision"] == b["decision"] == "ALLOW"


# ---------------------------------------------------------------- uploads
def demo_df():
    return pd.read_csv(SAMPLES / "policy_rules_demo.csv")


def upload(h, name, data):
    return c.post("/api/datasets", headers=h, files={"file": (name, data)})


def docx_bytes(df):
    d = Document()
    t = d.add_table(rows=1, cols=len(df.columns))
    for i, col in enumerate(df.columns):
        t.rows[0].cells[i].text = col
    for rec in df.itertuples(index=False):
        cells = t.add_row().cells
        for i, v in enumerate(rec):
            cells[i].text = str(v)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json", "tsv", "docx"])
def test_upload_every_supported_format(h, fmt):
    df = demo_df()
    if fmt == "csv": data = df.to_csv(index=False).encode()
    elif fmt == "tsv": data = df.to_csv(index=False, sep="\t").encode()
    elif fmt == "json": data = df.to_json(orient="records").encode()
    elif fmt == "xlsx":
        b = io.BytesIO(); df.to_excel(b, index=False); data = b.getvalue()
    else: data = docx_bytes(df)
    r = upload(h, f"demo.{fmt}", data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows"] == len(df) and body["counts"].get("BLOCK", 0) >= 2 and body["flagged"] >= 8
    c.delete(f"/api/datasets/{body['dataset']['id']}", headers=h)


@pytest.mark.parametrize("name,kind,code", [("virus.exe", "exe", 400), ("empty.csv", "empty", 400), ("bad.xlsx", "junk", 400)])
def test_bad_uploads_rejected(h, name, kind, code):
    data = {"exe": b"MZ", "empty": b"", "junk": b"not an excel file"}[kind]
    assert upload(h, name, data).status_code == code


def test_upload_size_limit(h, monkeypatch):
    monkeypatch.setattr(datasets, "MAX_BYTES", 10)
    r = upload(h, "big.csv", b"x" * 11)
    assert r.status_code == 413 and "300 MB" in r.text


def test_dataset_parser_has_no_row_count_limit():
    data = ("value\n" + "1\n" * 20_001).encode()
    assert len(datasets.parse("large.csv", data)) == 20_001


def test_unknown_category_is_adapted(h):
    df = demo_df(); df.loc[0, "merchant_category"] = "food"
    r = upload(h, "cat.csv", df.to_csv(index=False).encode())
    assert r.status_code == 200
    c.delete(f"/api/datasets/{r.json()['dataset']['id']}", headers=h)


def test_upload_adapts_common_column_names_and_unknown_extension(h):
    data = b"transaction_amount,transaction_time,user_id,payee,type\n500,2026-08-01 10:00:00,U1,R1,transfer\n85000,2026-08-01 10:10:00,U1,R1,payment\n"
    r = upload(h, "transactions.dat", data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows"] == 2 and body["counts"].get("STEP_UP_AUTH", 0) >= 1
    c.delete(f"/api/datasets/{body['dataset']['id']}", headers=h)


def test_upload_accepts_arbitrary_columns_with_neutral_defaults(h):
    r = upload(h, "custom.csv", b"customer_note,reference\nnormal payment,A1\nreview payment,A2\n")
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 2
    c.delete(f"/api/datasets/{r.json()['dataset']['id']}", headers=h)


def test_path_traversal_filename_is_harmless(h):
    r = upload(h, "../../evil.csv", demo_df().to_csv(index=False).encode())
    assert r.status_code == 200
    ds = r.json()["dataset"]
    assert c.get(f"/api/datasets/{ds['id']}/file", headers=h).status_code == 200
    c.delete(f"/api/datasets/{ds['id']}", headers=h)


def test_dataset_isolation_between_users(h):
    ds = upload(h, "mine.csv", demo_df().to_csv(index=False).encode()).json()["dataset"]
    other = make_user()
    assert ds["id"] not in [d["id"] for d in c.get("/api/datasets", headers=other).json()]
    assert c.get(f"/api/datasets/{ds['id']}/file", headers=other).status_code == 404
    assert c.delete(f"/api/datasets/{ds['id']}", headers=other).status_code == 404
    assert c.post(f"/api/datasets/{ds['id']}/rescore", headers=other).status_code == 404
    assert c.delete(f"/api/datasets/{ds['id']}", headers=h).status_code == 200


def test_rescore_replaces_old_results(h):
    ds = upload(h, "rs.csv", demo_df().to_csv(index=False).encode()).json()["dataset"]
    n1 = len(c.get(f"/api/export/flagged?format=json&dataset_id={ds['id']}", headers=h).json())
    assert c.post(f"/api/datasets/{ds['id']}/rescore", headers=h).status_code == 200
    n2 = len(c.get(f"/api/export/flagged?format=json&dataset_id={ds['id']}", headers=h).json())
    assert n1 == n2 > 0
    c.delete(f"/api/datasets/{ds['id']}", headers=h)


# ---------------------------------------------------------------- exports
@pytest.fixture(scope="module")
def dataset(h):
    ds = upload(h, "export.csv", demo_df().to_csv(index=False).encode()).json()["dataset"]
    yield ds
    c.delete(f"/api/datasets/{ds['id']}", headers=h)


def export(h, ds, fmt, **q):
    params = {"format": fmt, "dataset_id": ds["id"], **q}
    return c.get("/api/export/flagged", headers=h, params=params)


def test_export_xlsx_is_readable(h, dataset):
    r = export(h, dataset, "xlsx")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row > 5 and ws.max_column == 6
    assert ws.cell(1, 1).value == "Reference" and ws.cell(1, 4).value == "Risk %"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == ws.dimensions


def test_export_docx_is_readable(h, dataset):
    d = Document(io.BytesIO(export(h, dataset, "docx").content))
    assert len(d.tables[0].rows) > 5 and d.tables[0].rows[0].cells[0].text == "Reference"


def test_export_pdf_csv_json(h, dataset):
    assert export(h, dataset, "pdf").content[:5] == b"%PDF-"
    csv = export(h, dataset, "csv").content.decode("utf-8-sig")
    assert csv.splitlines()[0].startswith("Reference,Transaction time (UTC),Amount (INR),Risk %")
    rows = export(h, dataset, "json").json()
    assert rows and all(x["decision"] != "ALLOW" for x in rows)
    assert set(rows[0]) == {"reference", "time_utc", "amount_inr", "risk_percent", "decision", "why_flagged"}
    assert rows[0]["decision"] == "BLOCK"
    assert rows == sorted(rows, key=lambda x: ({"BLOCK": 0, "STEP_UP_AUTH": 1}[x["decision"]], -x["risk_percent"]))


def test_export_filters(h, dataset):
    high = export(h, dataset, "json", decision="BLOCK").json()
    assert high and all(x["decision"] == "BLOCK" and x["amount_inr"] >= 100_000 for x in high)
    mod = export(h, dataset, "json", decision="STEP_UP_AUTH").json()
    assert mod and all(x["decision"] == "STEP_UP_AUTH" for x in mod)
    assert export(h, dataset, "exe").status_code == 400
    assert export(h, dataset, "csv", status="confirmed_fraud").status_code == 404


# ---------------------------------------------------------------- alerts, simulation, overview
def test_alert_workflow(h):
    r = score(h, amount_inr=150_000)
    assert c.patch(f"/api/transactions/{r['id']}", headers=h, json={"status": "confirmed_fraud"}).json()["status"] == "confirmed_fraud"
    assert r["id"] in [x["id"] for x in c.get("/api/transactions?status=confirmed_fraud", headers=h).json()]
    assert c.patch(f"/api/transactions/{r['id']}", headers=h, json={"status": "bogus"}).status_code == 422
    assert c.patch("/api/transactions/99999999", headers=h, json={"status": "dismissed"}).status_code == 404


def test_overview_uses_scored_transactions(h):
    score(h, amount_inr=150_000)
    o = c.get("/api/overview", headers=h).json()
    assert o["total"] >= 1 and o["model"]["metrics"]["pr_auc"] > 0.5 and len(o["model"]["banks"]) == 5


# ---------------------------------------------------------------- privacy endpoint and training job
def test_epsilon_endpoint_monotonic(h):
    e = lambda z: c.get(f"/api/privacy/epsilon?noise={z}&rounds=15", headers=h).json()["epsilon"]
    assert e(2.0) < e(1.0) < e(0.5)
    assert c.get("/api/privacy/epsilon?noise=0&rounds=15", headers=h).json()["epsilon"] is None


def test_training_job_runs_and_blocks_duplicates(h):
    cfg = {"algorithm": "fedavg", "rounds": 2, "local_epochs": 1, "secure_agg": True, "promote": False, "full_comparison": False}
    first = c.post("/api/federated/train", headers=h, json=cfg)
    assert first.status_code == 200
    assert c.post("/api/federated/train", headers=h, json=cfg).status_code == 409
    for _ in range(120):
        st = c.get("/api/federated/status", headers=h).json()
        if st["status"] in ("done", "error"):
            break
        time.sleep(2)
    assert st["status"] == "done" and st["promoted"] is False and st["result"]["metrics"]["pr_auc"] > 0
    assert c.post("/api/federated/train", headers=h, json={**cfg, "rounds": 99}).status_code == 422
