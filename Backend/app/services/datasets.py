"""Parse uploaded files and adapt varied transaction schemas for fraud scoring."""

import io
import re
import uuid
from pathlib import Path

import pandas as pd

from app.config import ROOT

STORAGE = ROOT / "storage" / "datasets"
ALLOWED = {".csv", ".xlsx", ".xlsm", ".json", ".docx", ".txt", ".tsv"}
MAX_BYTES = 300 * 1024 * 1024

ALIASES = {
    "transaction_id": ("txn_id", "transaction_ref", "reference", "reference_id", "payment_id", "id"),
    "timestamp": ("transaction_time", "transaction_date", "date_time", "datetime", "created_at", "date", "time", "step"),
    "bank": ("bank_name", "institution", "provider"),
    "sender_id": ("sender", "payer", "payer_id", "customer_id", "user_id", "account_id", "source_account", "name_orig"),
    "receiver_id": ("receiver", "payee", "payee_id", "beneficiary", "beneficiary_id", "merchant_id", "destination_account", "name_dest"),
    "merchant_category": ("category", "merchant_type", "transaction_type", "payment_type", "type"),
    "amount_inr": ("amount", "transaction_amount", "payment_amount", "txn_amount", "amount_rs", "amount_rupees", "value", "amt"),
    "amount_to_average_ratio": ("amount_ratio", "amount_to_avg_ratio", "ratio_to_average", "amount_vs_average"),
    "transaction_hour": ("hour", "txn_hour"),
    "account_age_days": ("account_age", "customer_age_days", "days_since_account_opened"),
    "transactions_last_hour": ("txn_last_hour", "transactions_1h", "velocity_1h", "hourly_transactions"),
    "minutes_since_last_transaction": ("minutes_since_last", "time_since_last_transaction", "txn_gap_minutes"),
    "failed_login_attempts_24h": ("failed_logins", "failed_login_attempts", "login_failures_24h"),
    "upi_pin_failures_24h": ("pin_failures", "upi_pin_failures", "pin_failures_24h"),
    "is_new_device": ("new_device", "device_changed", "unknown_device"),
    "is_new_beneficiary": ("new_beneficiary", "new_payee", "first_payment_to_receiver"),
    "is_collect_request": ("collect_request", "is_pull_payment", "pull_request"),
    "is_weekend": ("weekend",),
    "is_fraud": ("fraud", "fraud_label", "fraud_flag", "isfraud", "label", "class"),
}

DEFAULTS = {
    "amount_inr": 0.0,
    "amount_to_average_ratio": 1.0,
    "transaction_hour": 14,
    "account_age_days": 800,
    "transactions_last_hour": 0,
    "minutes_since_last_transaction": 300.0,
    "failed_login_attempts_24h": 0,
    "upi_pin_failures_24h": 0,
    "is_new_device": 0,
    "is_new_beneficiary": 0,
    "is_collect_request": 0,
    "is_weekend": 0,
    "merchant_category": "p2p",
}


def parse(filename: str, raw: bytes) -> pd.DataFrame:
    ext = Path(filename).suffix.lower()
    if ext in (".csv", ".txt"):
        df = pd.read_csv(io.BytesIO(raw), sep=None, engine="python")
    elif ext == ".tsv":
        df = pd.read_csv(io.BytesIO(raw), sep="\t")
    elif ext in (".xlsx", ".xlsm"):
        df = pd.read_excel(io.BytesIO(raw), engine="openpyxl")
    elif ext == ".json":
        df = pd.read_json(io.BytesIO(raw))
    elif ext == ".docx":  # first table with a header row
        from docx import Document
        doc = Document(io.BytesIO(raw))
        if not doc.tables:
            raise ValueError("The Word document contains no tables.")
        t = doc.tables[0]
        data = [[c.text.strip() for c in r.cells] for r in t.rows]
        df = pd.DataFrame(data[1:], columns=data[0])
        for col in df.columns:                      # Word tables are text: convert numeric-looking columns
            try:
                df[col] = pd.to_numeric(df[col])
            except (ValueError, TypeError):
                pass
    else:
        # Unknown extensions are allowed when their contents are readable as a
        # delimited transaction table (for example .dat, .log or extensionless files).
        try:
            df = pd.read_csv(io.BytesIO(raw), sep=None, engine="python")
        except Exception as exc:
            raise ValueError("The file does not contain readable tabular transaction data.") from exc
    df.columns = [re.sub(r"[^a-z0-9]+", "_", str(c).strip().lower()).strip("_") for c in df.columns]
    df = df.dropna(how="all")
    if len(df) == 0:
        raise ValueError("The file has no data rows.")
    return df


def _as_flag(series: pd.Series) -> pd.Series:
    truthy = {"1", "true", "yes", "y", "new", "fraud", "fraudulent", "suspicious"}
    return series.map(lambda value: 1 if str(value).strip().lower() in truthy else 0)


def _category(value) -> str:
    text = str(value).strip().lower()
    if any(word in text for word in ("invest", "stock", "fund", "security")):
        return "investment"
    if any(word in text for word in ("p2p", "upi", "transfer", "cash_out", "person", "peer")):
        return "p2p"
    return "merchant"


def normalize_transactions(df: pd.DataFrame) -> pd.DataFrame:
    """Return a model-ready frame without requiring one exact source schema.

    Known aliases are renamed, useful fields are derived, and unavailable risk
    signals receive neutral defaults. Original extra columns are retained.
    """
    out = df.copy()
    rename = {}
    for canonical, aliases in ALIASES.items():
        if canonical not in out.columns:
            match = next((name for name in aliases if name in out.columns), None)
            if match:
                rename[match] = canonical
    out = out.rename(columns=rename)

    if "amount_inr" not in out:
        candidates = [c for c in out.columns if any(key in c for key in ("amount", "value", "amt"))]
        if candidates:
            out["amount_inr"] = pd.to_numeric(out[candidates[0]], errors="coerce")

    if "timestamp" in out:
        parsed = pd.to_datetime(out["timestamp"], errors="coerce", utc=True)
        valid = parsed.notna()
        if valid.all():
            out["timestamp"] = parsed
            if "transaction_hour" not in out:
                out["transaction_hour"] = parsed.dt.hour
            if "is_weekend" not in out:
                out["is_weekend"] = (parsed.dt.weekday >= 5).astype(int)
        else:
            out = out.drop(columns="timestamp")

    if "amount_to_average_ratio" not in out and "amount_inr" in out:
        amount = pd.to_numeric(out["amount_inr"], errors="coerce").clip(lower=0)
        if "sender_id" in out:
            average = amount.groupby(out["sender_id"].astype(str)).transform("median")
        else:
            median = amount[amount > 0].median()
            average = pd.Series(median if pd.notna(median) and median > 0 else 1.0, index=out.index)
        out["amount_to_average_ratio"] = (amount / average.replace(0, 1)).clip(0, 100).fillna(1.0)

    for column, default in DEFAULTS.items():
        if column not in out:
            out[column] = default

    flags = ("is_new_device", "is_new_beneficiary", "is_collect_request", "is_weekend")
    for column in flags:
        if not pd.api.types.is_numeric_dtype(out[column]):
            out[column] = _as_flag(out[column])
    numeric = [name for name, value in DEFAULTS.items() if not isinstance(value, str) and name not in flags]
    for column in numeric:
        out[column] = pd.to_numeric(out[column], errors="coerce").fillna(DEFAULTS[column])
    for column in flags:
        out[column] = pd.to_numeric(out[column], errors="coerce").fillna(0)
        out[column] = out[column].clip(0, 1).astype(int)

    out["amount_inr"] = out["amount_inr"].clip(lower=0)
    out["transaction_hour"] = out["transaction_hour"].clip(0, 23)
    out["merchant_category"] = out["merchant_category"].map(_category)
    if "is_fraud" in out:
        out["is_fraud"] = _as_flag(out["is_fraud"])
    if "transaction_id" not in out:
        out["transaction_id"] = [f"UPLOAD{i + 1}" for i in range(len(out))]
    if "bank" not in out:
        out["bank"] = "-"
    return out


def save_file(filename: str, raw: bytes) -> str:
    STORAGE.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}{Path(filename).suffix.lower()}"
    (STORAGE / stored).write_bytes(raw)
    return stored


def load_file(stored: str) -> bytes:
    return (STORAGE / stored).read_bytes()


def remove_file(stored: str):
    (STORAGE / stored).unlink(missing_ok=True)
